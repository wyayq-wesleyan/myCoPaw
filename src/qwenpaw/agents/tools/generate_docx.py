"""Tool to generate DOCX files from Markdown content."""

from __future__ import annotations

import os
import re
import uuid
from pathlib import Path

from agentscope.message import TextBlock, ToolResultState
from agentscope.tool import ToolChunk

from ...runtime.tool_registry import tool_descriptor


def _resolve_media_dir() -> Path:
    """Resolve the same media directory used by chat uploads.

    Mirrors ``ConsoleChannel._media_dir`` resolution (channel.py:112-117):
    ``{workspace_dir}/media`` when workspace_dir is set, else DEFAULT_MEDIA_DIR.
    """
    from ...app.channels.console.channel import DEFAULT_MEDIA_DIR

    workspace_dir = os.environ.get("QWENPAW_WORKSPACE_DIR", "").strip()
    media_dir = (
        Path(workspace_dir) / "media"
        if workspace_dir
        else DEFAULT_MEDIA_DIR
    )
    media_dir.mkdir(parents=True, exist_ok=True)
    return media_dir


def _markdown_to_docx(content: str, output_path: Path) -> None:
    """Convert markdown text to a DOCX file."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Microsoft YaHei"
    font.size = Pt(11)

    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Heading
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:], level=4)
        # Unordered list
        elif re.match(r"^[-*]\s+", line):
            text = re.sub(r"^[-*]\s+", "", line)
            doc.add_paragraph(text, style="List Bullet")
        # Ordered list
        elif re.match(r"^\d+\.\s+", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            doc.add_paragraph(text, style="List Number")
        # Table (detect | delimiter)
        elif line.startswith("|") and "|" in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            i -= 1  # Will be incremented at end of loop
            if len(table_lines) >= 2:
                rows = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.strip("|").split("|")]
                    # Skip separator row (---|---|---)
                    if all(re.match(r"^[-:]+$", c) for c in cells):
                        continue
                    rows.append(cells)
                if rows:
                    num_cols = max(len(r) for r in rows)
                    table = doc.add_table(rows=len(rows), cols=num_cols)
                    table.style = "Table Grid"
                    for ri, row_data in enumerate(rows):
                        for ci, cell_text in enumerate(row_data):
                            if ci < num_cols:
                                table.rows[ri].cells[ci].text = cell_text
        # Horizontal rule
        elif re.match(r"^[-*_]{3,}\s*$", line):
            doc.add_paragraph("─" * 50)
        # Empty line
        elif not line.strip():
            pass
        # Normal paragraph
        else:
            doc.add_paragraph(line)

        i += 1

    doc.save(str(output_path))


def _path_to_file_url(path: Path) -> str:
    """Convert a local file path to a ``file://`` URL (same as file_io.py:34)."""
    abs_path = str(path.resolve())
    if os.name == "nt":
        abs_path = abs_path.replace("\\", "/")
        return f"file:///{abs_path}"
    return f"file://{abs_path}"


@tool_descriptor(
    name="generate_docx",
    async_execution=False,
    tool_type="file",
    policy_name="GenerateDocx",
    default_policy="allow",
    policy_reason="Generates a DOCX file for user download",
    ui_description="将 Markdown 内容转换为 DOCX 文件，随后必须调用 send_file 发送给用户",
    ui_icon="📄",
)
def generate_docx(content: str, filename: str = "document") -> ToolChunk:
    """Generate a DOCX file from markdown content. Returns the file path —
    call ``send_file`` with the returned path to deliver the file to the user.

    Args:
        content: Markdown format document content
        filename: Output filename without extension
    """
    if not content or not content.strip():
        return ToolChunk(
            is_last=True,
            state=ToolResultState.ERROR,
            content=[TextBlock(type="text", text="内容不能为空")],
        )

    # Sanitize filename
    safe_name = re.sub(r'[<>:"/\\|?*]', "_", filename)[:100] or "document"
    stored_name = f"{uuid.uuid4().hex}_{safe_name}.docx"

    media_dir = _resolve_media_dir()
    output_path = media_dir / stored_name

    try:
        _markdown_to_docx(content, output_path)
    except ImportError:
        return ToolChunk(
            is_last=True,
            state=ToolResultState.ERROR,
            content=[TextBlock(type="text", text="python-docx 未安装，无法生成 DOCX 文件")],
        )
    except Exception as exc:
        return ToolChunk(
            is_last=True,
            state=ToolResultState.ERROR,
            content=[TextBlock(type="text", text=f"生成 DOCX 失败: {exc}")],
        )

    file_url = _path_to_file_url(output_path)
    return ToolChunk(
        is_last=True,
        state=ToolResultState.SUCCESS,
        content=[TextBlock(
            type="text",
            text=f"DOCX 文件已生成：{safe_name}.docx\n路径：{file_url}\n请调用 send_file 工具，参数 file_path 填上述路径，将文件发送给用户。",
        )],
    )